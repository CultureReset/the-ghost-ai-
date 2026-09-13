import os, re, sys, shutil, subprocess, datetime
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

class _Block:
    def find_module(self, name, path=None):
        if name == 'cryptography' or name.startswith('cryptography.'): return self
    def load_module(self, name): raise ImportError('blocked')
sys.meta_path.insert(0, _Block())

import wrap, mkdocx, h2md
from pypdf import PdfWriter

SCRATCH = os.path.dirname(os.path.abspath(__file__))
STAGE   = os.path.join(SCRATCH, 'docs')
CHROME  = '/opt/pw-browsers/chromium-1194/chrome-linux/chrome'

PARTS = [
 ("Part I",   "An Honest Read",      "honest-read.html",
  "What it is, what to bet on, the kiosk, what worries me, what I'd do, and the odds."),
 ("Part II",  "The Playbook",        "anextgent-playbook.html",
  "Thesis, products, architecture, data plane, economics, go-to-market, brand, IP, founder assets, risks."),
 ("Part III", "The Build Spec",      "build-spec.html",
  "Data placement, the business record, ingestion, capabilities, the app contract, surfaces, isolation, build order."),
 ("Part IV",  "The Build Plan",      "remote-control-build-plan.html",
  "Phases P0–P5 with done-when gates, the weekly ops loop, economics, the risk register, the first thirty days."),
 ("Part V",   "The Parts Catalog",   "parts-catalog.html",
  "Every open-source component by layer with a use / study / careful / skip verdict."),
 ("Part VI",  "The App Store Layer", "app-store.html",
  "Package format, index-not-store distribution, the trust ladder, the installer, nine deployable units, Grok Bot, Apple."),
]

def slug(s):
    return re.sub(r'[^a-z0-9]+', '-', s.lower()).strip('-')

def chrome_pdf(src, out):
    subprocess.run([CHROME, '--headless', '--disable-gpu', '--no-sandbox',
                    '--hide-scrollbars', '--virtual-time-budget=25000',
                    '--print-to-pdf-no-header', '--print-to-pdf=' + out, src],
                   cwd=STAGE, capture_output=True, timeout=240)
    ok = os.path.exists(out) and os.path.getsize(out) > 3000
    return ok

def main():
    os.makedirs(STAGE, exist_ok=True)
    for f in os.listdir(STAGE):
        if f.startswith('_') or f.endswith(('.pdf', '.docx', '.html', '.png')):
            os.remove(os.path.join(STAGE, f))

    made = []

    # ---- cover ----
    cover_html = os.path.join(STAGE, '00-cover.html')
    wrap.wrap(os.path.join(SCRATCH, 'cover.html'), cover_html)
    cover_pdf = os.path.join(STAGE, '_cover.pdf')
    assert chrome_pdf(cover_html, cover_pdf), 'cover pdf failed'

    part_pdfs = []
    for i, (pno, title, src, blurb) in enumerate(PARTS, 1):
        base = '%02d-%s' % (i, slug(title))
        html_out = os.path.join(STAGE, base + '.html')
        pdf_out  = os.path.join(STAGE, base + '.pdf')
        docx_out = os.path.join(STAGE, base + '.docx')

        wrap.wrap(os.path.join(SCRATCH, src), html_out, title=title)
        assert chrome_pdf(html_out, pdf_out), 'pdf failed: ' + base
        part_pdfs.append(pdf_out)

        doc = mkdocx.base_doc()
        mkdocx.part_title(doc, pno, title, blurb)
        mkdocx.render(doc, mkdocx.parse(os.path.join(SCRATCH, src)))
        doc.save(docx_out)

        made.append((base, os.path.getsize(pdf_out), os.path.getsize(docx_out)))
        print('%-34s pdf %7d  docx %6d' % (base, os.path.getsize(pdf_out),
                                           os.path.getsize(docx_out)))

    # ---- combined PDF, with bookmarks ----
    from pypdf import PdfReader
    w = PdfWriter()
    w.append(cover_pdf)
    page = len(PdfReader(cover_pdf).pages)
    w.add_outline_item('Cover & Contents', 0)
    for (pno, title, _s, _b), pdf in zip(PARTS, part_pdfs):
        w.append(pdf)
        w.add_outline_item('%s — %s' % (pno, title), page)
        page += len(PdfReader(pdf).pages)
    w.add_metadata({
        '/Title': 'A NEXT GENT — Complete Working Record',
        '/Subject': 'A Linux computer made simple, and the business around it.',
        '/Keywords': 'A NEXT GENT, remote control, small business, open source, bootc',
    })
    if hasattr(w, 'compress_identical_objects'):
        w.compress_identical_objects()
    combined = os.path.join(STAGE, 'A-NEXT-GENT-Complete-Record.pdf')
    w.write(combined); w.close()
    print('combined pdf %d bytes, %d pages' % (os.path.getsize(combined), page))

    # ---- combined DOCX ----
    doc = mkdocx.base_doc()
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph(); r = p.add_run('WORKING RECORD  ·  REVISION 3')
    r.font.name = mkdocx.MONO; r.font.size = Pt(8); r.bold = True
    r.font.color.rgb = mkdocx.ACC
    p = doc.add_paragraph(); r = p.add_run('A NEXT GENT')
    r.font.name = mkdocx.SANS; r.font.size = Pt(40); r.bold = True
    p = doc.add_paragraph()
    r = p.add_run('A Linux computer made simple — and the business around it.')
    r.font.name = mkdocx.SERIF; r.font.size = Pt(13); r.italic = True
    r.font.color.rgb = mkdocx.INK2
    mkdocx.bottom_rule(p)
    p = doc.add_paragraph(); r = p.add_run('Contents')
    r.font.name = mkdocx.SANS; r.font.size = Pt(13); r.bold = True
    for i, (pno, title, src, blurb) in enumerate(PARTS, 1):
        cp = doc.add_paragraph()
        cp.paragraph_format.space_after = Pt(5)
        r = cp.add_run('%02d  ' % i)
        r.font.name = mkdocx.MONO; r.font.size = Pt(9); r.bold = True
        r.font.color.rgb = mkdocx.ACC
        r = cp.add_run(title + '  ')
        r.font.name = mkdocx.SANS; r.font.size = Pt(11); r.bold = True
        r = cp.add_run(blurb)
        r.font.name = mkdocx.SERIF; r.font.size = Pt(10); r.font.color.rgb = mkdocx.INK2

    for i, (pno, title, src, blurb) in enumerate(PARTS, 1):
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        mkdocx.part_title(doc, pno, title, blurb)
        mkdocx.render(doc, mkdocx.parse(os.path.join(SCRATCH, src)))
    cdocx = os.path.join(STAGE, 'A-NEXT-GENT-Complete-Record.docx')
    doc.save(cdocx)
    print('combined docx', os.path.getsize(cdocx))

    os.remove(cover_pdf)
    return made

if __name__ == '__main__':
    main()
